"""Génère le dossier d'exemple des tests : tous les formats explorés et les cas limites.

Usage manuel : ``uv run python -m tests.corpus.make_corpus /tmp/corpus``
"""

from __future__ import annotations

import io
import os
import sys
import unicodedata
import zipfile
from dataclasses import dataclass
from email.message import EmailMessage
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont

LONG_PDF_PAGES = 45
LONG_TEXT_CHARS = 65_000


# --- PDF écrits à la main ------------------------------------------------------------


@dataclass
class PdfPage:
    lines: list[str]
    image_jpeg: bytes | None = None
    image_size: tuple[int, int] = (0, 0)


def _pdf_text(text: str) -> bytes:
    raw = text.encode("cp1252", errors="replace")
    out = bytearray()
    for byte in raw:
        char = chr(byte)
        if char in "()\\":
            out += b"\\" + char.encode()
        elif byte < 32 or byte > 126:
            out += f"\\{byte:03o}".encode()
        else:
            out.append(byte)
    return bytes(out)


def write_pdf(path: Path, pages: list[PdfPage], width: int = 595, height: int = 842) -> None:
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")  # rempli plus bas
    pages_id = add(b"")
    font_id = add(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    kids: list[int] = []
    for page in pages:
        content = bytearray()
        resources = f"/Font << /F1 {font_id} 0 R >>".encode()
        if page.image_jpeg is not None:
            img_w, img_h = page.image_size
            image_id = add(
                b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceRGB"
                b" /BitsPerComponent 8 /Filter /DCTDecode /Length %d >>\nstream\n"
                % (img_w, img_h, len(page.image_jpeg))
                + page.image_jpeg
                + b"\nendstream"
            )
            resources += f" /XObject << /Im1 {image_id} 0 R >>".encode()
            content += f"q {width} 0 0 {height} 0 0 cm /Im1 Do Q\n".encode()
        if page.lines:
            content += b"BT /F1 11 Tf 14 TL 50 800 Td\n"
            for line in page.lines:
                content += b"(" + _pdf_text(line) + b") Tj T*\n"
            content += b"ET\n"
        stream_id = add(
            b"<< /Length %d >>\nstream\n" % len(content) + bytes(content) + b"\nendstream"
        )
        kids.append(
            add(
                f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 {width} {height}]"
                f" /Resources << {resources.decode()} >> /Contents {stream_id} 0 R >>".encode()
            )
        )
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode()
    kids_ref = " ".join(f"{k} 0 R" for k in kids)
    objects[pages_id - 1] = f"<< /Type /Pages /Kids [{kids_ref}] /Count {len(kids)} >>".encode()

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\n"
    out += f"{trailer}startxref\n{xref}\n%%EOF\n".encode()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(out))


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for candidate in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
    ):
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size=size)


def text_image(lines: list[str], size: tuple[int, int] = (1654, 2339)) -> Image.Image:
    """Page blanche portant du texte, comme un scan à 200 dpi."""
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    font = _font(56)
    y = 150
    for line in lines:
        draw.text((120, y), line, fill="black", font=font)
        y += 110
    return image


def _jpeg(image: Image.Image) -> bytes:
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=90)
    return buffer.getvalue()


# --- corpus --------------------------------------------------------------------------


@dataclass(frozen=True)
class Corpus:
    root: Path
    nfd_name: str  # nom stocké en NFD sur le disque
    nfc_rel: str  # même chemin en NFC, tel que le voit le modèle


SCAN_LINES = ["LOYER MENSUEL", "Montant 850 euros", "Quittance de mars"]
IMAGE_LINES = ["FACTURE ELECTRICITE", "Total 245 euros"]


def build_corpus(root: Path) -> Corpus:
    root.mkdir(parents=True, exist_ok=True)

    # Fichiers texte
    (root / "notes.txt").write_text(
        "Réunion du 3 mars.\nL'échéance du bail est fixée au 30 juin.\n", encoding="utf-8"
    )
    (root / "ancien.txt").write_bytes(
        "Loyer annuel : 12 000 € — échéance trimestrielle\r\n".encode("cp1252")
    )
    (root / "abime.txt").write_bytes(b"debut \x81\x8d\xff fin du texte avec quittance\n")
    (root / "donnees.csv").write_text(
        "date;libelle;montant\n2024-01-05;Loyer janvier;850\n2024-02-05;Loyer février;850\n",
        encoding="utf-8-sig",
    )
    (root / "README.md").write_text("# Dossier\n\nDocuments de test.\n", encoding="utf-8")
    nfd_name = unicodedata.normalize("NFD", "été_résumé.txt")
    (root / nfd_name).write_text("Résumé estival : piscine municipale.\n", encoding="utf-8")

    # PDF
    contrats = root / "contrats"
    write_pdf(
        contrats / "bail_2022.pdf",
        [
            PdfPage(
                ["CONTRAT DE BAIL", "Entre les soussignés, le bailleur et le locataire."]
                + [
                    f"Article {i} : dispositions générales du contrat de location."
                    for i in range(1, 20)
                ]
            ),
            PdfPage(
                [
                    "Article 20 : clause de résiliation anticipée.",
                    "Le locataire peut résilier le bail avec un préavis de trois mois.",
                    "Le loyer annuel est de 12 000 euros hors charges.",
                ]
                + [f"Ligne de remplissage numéro {i} pour la deuxième page." for i in range(20)]
            ),
            PdfPage(
                [
                    "Article 30 : date d'échéance du bail : 30 juin 2025.",
                    "Fait à Toulon, en deux exemplaires.",
                ]
                + [f"Mention finale numéro {i} du contrat signé." for i in range(20)]
            ),
        ],
    )
    scan = text_image(SCAN_LINES)
    write_pdf(
        contrats / "scan_quittance.pdf",
        [
            PdfPage(
                ["Scanned by Archipelle Test — document numérisé le 12 mars 2024 par le courrier"],
                _jpeg(scan),
                scan.size,
            ),
            PdfPage([], _jpeg(text_image(["Seconde page", "LOYER MENSUEL avril"])), scan.size),
        ],
    )
    write_pdf(
        contrats / "mixte.pdf",
        [
            PdfPage(
                [
                    f"Page native du rapport mixte, paragraphe {i} sur le chauffage."
                    for i in range(12)
                ]
            ),
            PdfPage([], _jpeg(scan), scan.size),
        ],
    )
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    writer.encrypt("secret", algorithm="RC4-128")
    with (contrats / "protege.pdf").open("wb") as handle:
        writer.write(handle)
    (contrats / "corrompu.pdf").write_bytes(b"%PDF-1.4\nceci n'est pas un vrai pdf\n")

    # Bureautique
    bureau = root / "bureau"
    bureau.mkdir()
    import docx

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "En-tête confidentiel"
    document.sections[0].footer.paragraphs[0].text = "Pied de page société"
    document.add_heading("Rapport annuel", level=1)
    document.add_paragraph("Premier paragraphe du rapport.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Échéance"
    table.cell(0, 1).text = "30/06/2025"
    table.cell(1, 0).text = "Montant"
    table.cell(1, 1).text = "4 500 €"
    document.add_paragraph("Paragraphe après le tableau.")
    document.save(str(bureau / "rapport.docx"))
    (bureau / "protege.docx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 512)

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Budget"
    sheet.append(["Poste", "Montant"])
    sheet.append(["Chauffage", 1200.5])
    sheet.append(["Électricité", 800])
    sheet.append(["Total", "=SUM(B2:B3)"])
    second = workbook.create_sheet("Notes")
    second.append(["Relevé du compteur"])
    workbook.save(bureau / "budget.xlsx")

    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    shapes: Any = slide.shapes
    shapes.title.text = "Présentation du projet"
    placeholder: Any = slide.placeholders[1]
    placeholder.text = "Calendrier et budget prévisionnel"
    notes: Any = slide.notes_slide.notes_text_frame
    notes.text = "Note orale : insister sur l'échéance."
    second_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    rows = second_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    rows.cell(0, 0).text = "Phase"
    rows.cell(0, 1).text = "Coût"
    rows.cell(1, 0).text = "Étude"
    rows.cell(1, 1).text = "3 000 €"
    presentation.save(str(bureau / "presentation.pptx"))

    # E-mail et HTML
    message = EmailMessage()
    message["From"] = "Agence <agence@example.com>"
    message["To"] = "locataire@example.com"
    message["Subject"] = "Révision du loyer"
    message["Date"] = "Mon, 04 Mar 2024 10:00:00 +0100"
    message.set_content("Bonjour,\nLe loyer sera révisé au 1er avril.\nCordialement")
    message.add_attachment(
        b"%PDF-1.4", maintype="application", subtype="pdf", filename="avenant.pdf"
    )
    mails = root / "mails"
    mails.mkdir()
    (mails / "message.eml").write_bytes(bytes(message))
    web = root / "web"
    web.mkdir()
    (web / "page.html").write_text(
        "<html><head><title>Guide du locataire</title><style>p{color:red}</style>"
        "<script>var cache = 'motsecret';</script></head><body><h1>Guide</h1>"
        "<p>Les charges &amp; le d&eacute;p&ocirc;t de garantie.</p>"
        "<table><tr><td>Caution</td><td>1 mois</td></tr></table></body></html>",
        encoding="utf-8",
    )

    # Images
    images = root / "images"
    images.mkdir()
    text_image(IMAGE_LINES, (1400, 500)).save(images / "facture.png")
    Image.new("RGB", (400, 300), "white").save(images / "vide.png")
    (images / "casse.jpg").write_bytes(b"pas une image")

    # Formats non explorés et éléments exclus
    archives = root / "archives"
    archives.mkdir()
    with zipfile.ZipFile(archives / "lot.zip", "w") as archive:
        archive.writestr("dedans.txt", "échéance cachée dans une archive")
    (archives / "ancien.doc").write_bytes(b"\xd0\xcf\x11\xe0 contenu doc")
    (root / ".cache_cachee").mkdir()
    (root / ".cache_cachee" / "secret.txt").write_text("échéance secrète", encoding="utf-8")
    (root / ".git").mkdir()
    (root / ".git" / "config").write_text("[core]", encoding="utf-8")
    (root / "node_modules").mkdir()
    (root / "node_modules" / "paquet.txt").write_text("échéance npm", encoding="utf-8")
    (bureau / "~$rapport.docx").write_bytes(b"verrou")
    (root / "Thumbs.db").write_bytes(b"\0")
    (root / ".DS_Store").write_bytes(b"\0")

    # Gros fichiers
    gros = root / "gros"
    gros.mkdir()
    sentence = "Ligne de texte volumineux pour tester la troncature des lectures. "
    body = (sentence * (LONG_TEXT_CHARS // len(sentence) + 1))[:LONG_TEXT_CHARS]
    (gros / "long.txt").write_text(body + "\nMOT FINAL introuvable ailleurs\n", encoding="utf-8")
    write_pdf(
        gros / "long.pdf",
        [
            PdfPage(
                [f"Page {n} ligne {i} : texte du long document de test." for i in range(40)]
                + (["Mention spéciale : indexation du loyer."] if n == 33 else [])
            )
            for n in range(1, LONG_PDF_PAGES + 1)
        ],
    )

    # Dates de modification maîtrisées
    old = 1_577_836_800  # 2020-01-01
    os.utime(root / "ancien.txt", (old, old))
    return Corpus(root=root, nfd_name=nfd_name, nfc_rel=unicodedata.normalize("NFC", nfd_name))


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/corpus/generated")
    corpus = build_corpus(target)
    print(f"Corpus généré dans {corpus.root}")
