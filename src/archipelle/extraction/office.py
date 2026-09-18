"""Formats bureautiques : .docx, .xlsx, .pptx (CDC §3)."""

from __future__ import annotations

import datetime as dt
from collections.abc import Iterable, Iterator
from itertools import zip_longest
from typing import Any

from archipelle.core.i18n import t
from archipelle.extraction.errors import check_office_container, os_errors
from archipelle.extraction.types import (
    MAX_DOCUMENT_CHARS,
    DocumentNote,
    ErrorCode,
    ExtractionError,
    ExtractionMethod,
    NoteEntry,
    WholeDocument,
)


class _TextBuilder:
    """Accumule des lignes en respectant la taille maximale d'un document."""

    def __init__(self) -> None:
        self.parts: list[str] = []
        self.size = 0
        self.truncated = False

    def add(self, line: str) -> bool:
        if self.truncated:
            return False
        if self.size + len(line) + 1 > MAX_DOCUMENT_CHARS:
            self.truncated = True
            return False
        self.parts.append(line)
        self.size += len(line) + 1
        return True

    def result(self, notes: list[NoteEntry]) -> WholeDocument:
        if self.truncated:
            notes.append(NoteEntry(DocumentNote.TRUNCATED))
        text = "\n".join(self.parts).strip()
        if not text:
            notes.append(NoteEntry(DocumentNote.NO_TEXT))
        return WholeDocument(text=text, method=ExtractionMethod.NATIVE, notes=notes)


def _row_text(cells: Iterable[str]) -> str:
    """Cellules d'une ligne de tableau, en retirant les répétitions des cellules fusionnées."""
    values: list[str] = []
    for cell in cells:
        cleaned = " ".join(cell.split())
        if not values or cleaned != values[-1]:
            values.append(cleaned)
    return " | ".join(values).strip(" |")


# --- docx ----------------------------------------------------------------------------


def _docx_blocks(container: Any) -> Iterator[str]:
    from docx.table import Table

    for block in container.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                line = _row_text(cell.text for cell in row.cells)
                if line:
                    yield line
        else:
            text = block.text.strip()
            if text:
                yield text


def extract_docx(path: str) -> WholeDocument:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    check_office_container(path)
    with os_errors():
        try:
            document = docx.Document(path)
        except (PackageNotFoundError, KeyError, ValueError) as exc:
            raise ExtractionError(ErrorCode.CORRUPT, str(exc)) from exc

    headers: list[str] = []
    footers: list[str] = []
    for section in document.sections:
        for part, target in ((section.header, headers), (section.footer, footers)):
            if part.is_linked_to_previous:
                continue
            text = "\n".join(_docx_blocks(part))
            if text and text not in target:
                target.append(text)

    builder = _TextBuilder()
    for text in headers:
        builder.add(t("extraction.docx.header"))
        builder.add(text)
    for line in _docx_blocks(document):
        if not builder.add(line):
            break
    for text in footers:
        builder.add(t("extraction.docx.footer"))
        builder.add(text)
    return builder.result([])


# --- xlsx ----------------------------------------------------------------------------


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(" ")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return " ".join(str(value).split())


def _is_formula(value: Any) -> bool:
    return (isinstance(value, str) and value.startswith("=")) or type(value).__name__ in {
        "ArrayFormula",
        "DataTableFormula",
    }


def extract_xlsx(path: str) -> WholeDocument:
    import openpyxl

    check_office_container(path)
    with os_errors():
        try:
            values_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        except (KeyError, ValueError, TypeError) as exc:
            raise ExtractionError(ErrorCode.CORRUPT, str(exc)) from exc

    builder = _TextBuilder()
    missing = 0
    try:
        for sheet_name in values_book.sheetnames:
            if builder.truncated:
                break
            builder.add(t("extraction.xlsx.sheet", name=sheet_name))
            value_rows = values_book[sheet_name].iter_rows(values_only=True)
            formula_rows = formula_book[sheet_name].iter_rows(values_only=True)
            for values, formulas in zip_longest(value_rows, formula_rows, fillvalue=()):
                cells: list[str] = []
                for value, formula in zip_longest(values, formulas):
                    if value is None and _is_formula(formula):
                        missing += 1
                    cells.append(_cell_text(value))
                while cells and not cells[-1]:
                    cells.pop()
                if cells and not builder.add("\t".join(cells)):
                    break
    finally:
        values_book.close()
        formula_book.close()

    notes = [NoteEntry(DocumentNote.XLSX_MISSING_VALUES, str(missing))] if missing else []
    return builder.result(notes)


# --- pptx ----------------------------------------------------------------------------


def _pptx_shape_lines(shapes: Iterable[Any]) -> Iterator[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_shape_lines(shape.shapes)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                yield text
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                line = _row_text(cell.text for cell in row.cells)
                if line:
                    yield line


def extract_pptx(path: str) -> WholeDocument:
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    check_office_container(path)
    with os_errors():
        try:
            presentation = Presentation(path)
        except (PackageNotFoundError, KeyError, ValueError) as exc:
            raise ExtractionError(ErrorCode.CORRUPT, str(exc)) from exc

    builder = _TextBuilder()
    for number, slide in enumerate(presentation.slides, start=1):
        builder.add(t("extraction.pptx.slide", number=number))
        for line in _pptx_shape_lines(slide.shapes):
            builder.add(line)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""
            if notes:
                builder.add(t("extraction.pptx.notes"))
                builder.add(notes)
        if builder.truncated:
            break
    return builder.result([])
